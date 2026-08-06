"""Parse CSV and XLSX uploads into column headers and row dicts."""

import csv
import io
from typing import Iterator

from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}


def file_extension(filename: str) -> str:
    name = (filename or "").lower().strip()
    if name.endswith(".xlsm"):
        return ".xlsm"
    if name.endswith(".xlsx"):
        return ".xlsx"
    if name.endswith(".csv"):
        return ".csv"
    return ""


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_csv(raw: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV must be UTF-8 encoded") from exc

    if not text.strip():
        raise ValueError("File is empty")

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("File has no header row")

    columns = [c for c in reader.fieldnames if c is not None]
    rows = []
    for row in reader:
        rows.append({col: (row.get(col) or "").strip() for col in columns})
    return columns, rows


def parse_xlsx(raw: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Invalid or corrupted XLSX file") from exc

    ws = wb.active
    if ws is None:
        raise ValueError("XLSX workbook has no sheets")

    row_iter: Iterator = ws.iter_rows(values_only=True)
    try:
        header = next(row_iter)
    except StopIteration:
        raise ValueError("File is empty")

    columns = [_cell_str(c) or f"column_{i + 1}" for i, c in enumerate(header)]
    if not any(columns):
        raise ValueError("File has no header row")

    rows: list[dict[str, str]] = []
    for values in row_iter:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        row_dict = {}
        for i, col in enumerate(columns):
            row_dict[col] = _cell_str(values[i]) if i < len(values) else ""
        rows.append(row_dict)

    wb.close()
    return columns, rows


def parse_spreadsheet(raw: bytes, filename: str) -> tuple[list[str], list[dict[str, str]]]:
    ext = file_extension(filename)
    if ext == ".csv":
        return parse_csv(raw)
    if ext in (".xlsx", ".xlsm"):
        return parse_xlsx(raw)
    raise ValueError("Unsupported file type. Upload CSV or XLSX.")
